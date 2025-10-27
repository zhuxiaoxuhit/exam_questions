#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将Word格式的考题转换为XLS格式

输入格式（Word文档）：
A-B-A-001  B  3  5
{A}题干内容
（A）选项A
（B）选项B
（C）选项C
（D）选项D
（E）选项E（仅多选题）
{B}答案

输出格式（XLS文件）：
11列表格，包含鉴定点代码、题型、题干、4-5个选项、答案、难度、一致性
"""

import os
import sys
import re
import xlwt
from docx import Document


class DocxToXlsConverter:
    """Word考题转XLS转换器"""

    def __init__(self):
        """初始化转换器"""
        pass

    def parse_docx(self, docx_filepath):
        """
        解析Word文档，提取题目信息

        Args:
            docx_filepath: Word文档路径

        Returns:
            list: 题目列表，每个题目是一个字典
        """
        doc = Document(docx_filepath)
        questions = []
        current_question = None

        for para in doc.paragraphs:
            text = para.text.strip()

            if not text:
                continue

            # 匹配元数据行：A-B-A-001  B  3  5
            meta_match = re.match(r'^([A-Z]-[A-Z]-[A-Z]-\d+)\s+([BCD])\s+(\d+)\s+(\d+)$', text)
            if meta_match:
                # 如果已有题目，保存它
                if current_question:
                    questions.append(current_question)

                # 创建新题目
                current_question = {
                    '鉴定点代码': meta_match.group(1),
                    '题型': meta_match.group(2),
                    '难度': meta_match.group(3),
                    '一致性': meta_match.group(4),
                    '题干': '',
                    '选项A': '',
                    '选项B': '',
                    '选项C': '',
                    '选项D': '',
                    '选项E': '',
                    '答案': ''
                }
                continue

            # 匹配题干：{A}题干内容
            stem_match = re.match(r'^\{A\}(.+)$', text)
            if stem_match and current_question:
                current_question['题干'] = stem_match.group(1)
                continue

            # 匹配选项：（A）选项内容
            option_match = re.match(r'^（([A-E])）(.+)$', text)
            if option_match and current_question:
                option_letter = option_match.group(1)
                option_content = option_match.group(2)
                current_question[f'选项{option_letter}'] = option_content
                continue

            # 匹配答案：{B}答案
            answer_match = re.match(r'^\{B\}(.+)$', text)
            if answer_match and current_question:
                current_question['答案'] = answer_match.group(1)
                continue

        # 保存最后一个题目
        if current_question:
            questions.append(current_question)

        return questions

    def save_to_xls(self, questions, xls_filepath):
        """
        将题目保存为XLS文件

        Args:
            questions: 题目列表
            xls_filepath: XLS文件路径
        """
        # 创建工作簿
        wb = xlwt.Workbook(encoding='utf-8')
        ws = wb.add_sheet('题目')

        # 设置宋体样式
        style = xlwt.XFStyle()
        font = xlwt.Font()
        font.name = '宋体'
        font.height = 220  # 11号字体
        style.font = font

        # 写入表头
        headers = [
            '鉴定点代码', '题目类型代码', '试题(题干)', '选项A', '选项B',
            '选项C', '选项D', '选项E', '答案', '难度代码', '一致性代码'
        ]

        for col_idx, header in enumerate(headers):
            ws.write(0, col_idx, header, style)

        # 写入题目数据
        for row_idx, question in enumerate(questions, start=1):
            ws.write(row_idx, 0, question['鉴定点代码'], style)
            ws.write(row_idx, 1, question['题型'], style)
            ws.write(row_idx, 2, question['题干'], style)
            ws.write(row_idx, 3, question['选项A'], style)
            ws.write(row_idx, 4, question['选项B'], style)
            ws.write(row_idx, 5, question['选项C'], style)
            ws.write(row_idx, 6, question['选项D'], style)
            ws.write(row_idx, 7, question['选项E'], style)
            ws.write(row_idx, 8, question['答案'], style)
            ws.write(row_idx, 9, question['难度'], style)
            ws.write(row_idx, 10, question['一致性'], style)

        # 保存文件
        wb.save(xls_filepath)
        print(f"[OK] 已保存到: {xls_filepath}")

    def convert_file(self, docx_filepath, xls_filepath=None):
        """
        转换单个文件

        Args:
            docx_filepath: Word文档路径
            xls_filepath: XLS文件路径，如果为None则自动生成
        """
        if not os.path.exists(docx_filepath):
            print(f"❌ 错误: 文件不存在 - {docx_filepath}")
            return

        print(f"正在处理: {docx_filepath}")

        # 解析Word文档
        questions = self.parse_docx(docx_filepath)

        if not questions:
            print(f"⚠ 警告: 未找到任何题目")
            return

        print(f"[OK] 解析到 {len(questions)} 道题目")

        # 确定输出文件路径
        if xls_filepath is None:
            # 从 A-B-A-001.docx 转换为 考题A-B-A-001.xls
            basename = os.path.basename(docx_filepath)
            name_without_ext = os.path.splitext(basename)[0]
            xls_filename = f"考题{name_without_ext}.xls"
            xls_filepath = os.path.join(os.path.dirname(docx_filepath), xls_filename)

        # 保存为XLS
        self.save_to_xls(questions, xls_filepath)

    def convert_directory(self, directory):
        """
        转换目录中的所有docx文件

        Args:
            directory: 目录路径
        """
        if not os.path.exists(directory):
            print(f"❌ 错误: 目录不存在 - {directory}")
            return

        # 查找所有docx文件（排除临时文件）
        docx_files = []
        for filename in os.listdir(directory):
            if filename.endswith('.docx') and not filename.startswith('~$'):
                docx_files.append(os.path.join(directory, filename))

        if not docx_files:
            print(f"⚠ 警告: 目录中没有找到docx文件")
            return

        print(f"找到 {len(docx_files)} 个Word文档")
        print("="*60)

        # 转换每个文件
        for docx_file in sorted(docx_files):
            self.convert_file(docx_file)
            print()

        print("="*60)
        print(f"[OK] 完成! 共处理 {len(docx_files)} 个文件")


def main():
    """主函数"""
    converter = DocxToXlsConverter()

    if len(sys.argv) < 2:
        print("用法:")
        print("  python docx_to_xls.py <docx文件路径>        # 转换单个文件")
        print("  python docx_to_xls.py <目录路径>            # 转换目录中所有docx文件")
        print()
        print("示例:")
        print("  python docx_to_xls.py A-B-A-001.docx")
        print("  python docx_to_xls.py 考题")
        return

    input_path = sys.argv[1]

    # 判断是文件还是目录
    if os.path.isfile(input_path):
        # 转换单个文件
        if len(sys.argv) >= 3:
            xls_filepath = sys.argv[2]
        else:
            xls_filepath = None
        converter.convert_file(input_path, xls_filepath)

    elif os.path.isdir(input_path):
        # 转换目录中的所有文件
        converter.convert_directory(input_path)

    else:
        print(f"❌ 错误: 路径不存在 - {input_path}")


if __name__ == "__main__":
    main()
